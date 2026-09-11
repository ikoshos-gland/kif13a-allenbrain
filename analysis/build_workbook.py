"""Write the numeric tables from the analysis into a plain Excel workbook.

One sheet per topic, tables stacked with a blank row between them. Values only,
no formulas and no commentary. The first sheet carries the headline tables so
they are visible the moment the file opens.

Output: Kif13a_sonuclar.xlsx in the repository root.
"""
from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
OUT = ROOT / "Kif13a_sonuclar.xlsx"

FONT = "Arial"
HDR_FONT = Font(name=FONT, bold=True, size=10, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="44546A")
BODY = Font(name=FONT, size=10)
CAPTION = Font(name=FONT, bold=True, size=11)

F3 = "0.000"
F4 = "0.0000"
PCT = "0.0%"
INT = "#,##0"

wb = Workbook()
wb.remove(wb.active)


def put(ws, row, caption, df, fmt=None):
    """Write one captioned table. Returns the next free row."""
    ws.cell(row=row, column=1, value=caption).font = CAPTION
    row += 1
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=row, column=j, value=col)
        c.font, c.fill = HDR_FONT, HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for _, rec in df.iterrows():
        row += 1
        for j, col in enumerate(df.columns, start=1):
            v = rec[col]
            c = ws.cell(row=row, column=j, value=None if pd.isna(v) else v)
            c.font = BODY
            if fmt and col in fmt:
                c.number_format = fmt[col]
    return row + 2


def autofit(ws, cap=44):
    for col in ws.columns:
        letter, best = None, 0
        for c in col:
            if letter is None:
                letter = c.column_letter
            if c.value is not None:
                best = max(best, len(str(c.value)))
        if letter:
            ws.column_dimensions[letter].width = min(max(best + 2, 9), cap)


# ---------------------------------------------------------- shared source data
rank = pd.read_csv(RESULTS / "derived_wholebrain_subclass_ranking.csv")
rank.insert(0, "Sıra", range(1, len(rank) + 1))
rank = rank.set_index("subclass", drop=False)

TOP10 = [
    ("327 Oligo NN", "Oligo (olgun oligodendrosit)", "OPC-Oligo"),
    ("326 OPC NN", "OPC (öncül)", "OPC-Oligo"),
    ("145 MH Tac2 Glut", "MH Tac2 Glut", "Medial habenula"),
    ("298 PRP Gata3 Slc6a5 Gly-Gaba", "PRP Gata3 Slc6a5", "Medulla GABA"),
    ("251 NTS Dbh Glut", "NTS Dbh Glut", "Medulla"),
    ("294 MV Pax6 Gly-Gaba", "MV Pax6", "Medulla GABA"),
    ("292 MV Nkx6-1 Gly-Gaba", "MV Nkx6-1", "Medulla GABA"),
    ("287 MV-SPIV-PRP Dmbx1 Gly-Gaba", "MV-SPIV-PRP Dmbx1", "Medulla GABA"),
    ("283 PRP Otp Gly-Gaba", "PRP Otp", "Medulla GABA"),
    ("129 VMH Nr5a1 Glut", "VMH Nr5a1 Glut", "Hipotalamus"),
]
GLIA = [
    ("327 Oligo NN", "Olgun oligodendrosit"),
    ("326 OPC NN", "OPC"),
    ("319 Astro-TE NN", "Astrosit (telensefalik)"),
    ("334 Microglia NN", "Mikroglia"),
    ("333 Endo NN", "Endotel"),
    ("331 Peri NN", "Perisit"),
]
LOWEST = [
    ("040 OB Trdn Gaba", "OB Trdn Gaba"),
    ("042 OB-out Frmd7 Gaba", "OB-out Frmd7 Gaba"),
    ("045 OB-STR-CTX Inh IMN", "OB-STR-CTX Inh IMN"),
    ("041 OB-in Frmd7 Gaba", "OB-in Frmd7 Gaba"),
    ("047 Sncg Gaba", "Sncg Gaba (korteks)"),
    ("046 Vip Gaba", "Vip Gaba (korteks)"),
]
for key, *_ in TOP10 + GLIA + LOWEST:
    assert key in rank.index, "sıralamada yok: " + key

top10_df = pd.DataFrame({
    "Sıra": [int(rank.loc[k, "Sıra"]) for k, _, _ in TOP10],
    "Alt sınıf": [lab for _, lab, _ in TOP10],
    "Sınıf": [cls for _, _, cls in TOP10],
    "Hücre": [int(rank.loc[k, "n_cells"]) for k, _, _ in TOP10],
    "Ortalama (log2)": [float(rank.loc[k, "mean_Kif13a"]) for k, _, _ in TOP10],
})
glia_df = pd.DataFrame({
    "Hücre tipi": [lab for _, lab in GLIA],
    "Ortalama (log2)": [float(rank.loc[k, "mean_Kif13a"]) for k, _ in GLIA],
    "İfade eden oran": [float(rank.loc[k, "fraction_expressing"]) for k, _ in GLIA],
})
low_df = pd.DataFrame({
    "Alt sınıf": [lab for _, lab in LOWEST],
    "Ortalama (log2)": [float(rank.loc[k, "mean_Kif13a"]) for k, _ in LOWEST],
    "İfade eden oran": [float(rank.loc[k, "fraction_expressing"]) for k, _ in LOWEST],
})

donors = pd.read_csv(RESULTS / "Kif13a_AgingMouse_OPC_Oligo_donor_summary.csv")
age_test = (pd.read_csv(RESULTS / "derived_donor_level_age_test.csv")
            .set_index("subclass").reindex(["327 Oligo NN", "326 OPC NN"]).reset_index())
aging_main = pd.DataFrame({
    "Hücre tipi": ["Olgun oligodendrosit", "OPC (öncül)"],
    "Genç ortalama (log2)": age_test.mean_adult.values,
    "Yaşlı ortalama (log2)": age_test.mean_aged.values,
    "Fark (log2)": age_test.delta_log2.values,
    "Kat değişim": age_test.fold_change.values,
    "Mann-Whitney p": age_test.mannwhitney_p.values,
    "Welch t p": age_test.welch_t_p.values,
    "Cliff delta": age_test.cliffs_delta.values,
})
AGING_FMT = {"Genç ortalama (log2)": F3, "Yaşlı ortalama (log2)": F3,
             "Fark (log2)": "+0.000;-0.000", "Kat değişim": F3,
             "Mann-Whitney p": F4, "Welch t p": F4, "Cliff delta": "+0.000;-0.000"}

# ------------------------------------------------------------- 1. headline sheet
ws = wb.create_sheet("Özet tablolar")
r = put(ws, 1, "En yüksek 10 hücre tipi", top10_df, {"Hücre": INT, "Ortalama (log2)": F3})
r = put(ws, r, "Glia tipleri karşılaştırması", glia_df,
        {"Ortalama (log2)": F3, "İfade eden oran": "0.00"})
r = put(ws, r, "En düşük hücreler", low_df,
        {"Ortalama (log2)": F3, "İfade eden oran": "0.00"})
put(ws, r, "Genç ve yaşlı farelerde Kif13a, donör düzeyi", aging_main, AGING_FMT)
autofit(ws)

# ------------------------------------------------------------- 2. aging result
ws = wb.create_sheet("Yaşlanma sonucu")
r = put(ws, 1, "Genç ve yaşlı farelerde Kif13a, donör düzeyi", aging_main, AGING_FMT)

conv = pd.DataFrame({
    "Ölçüm": ["Oligodendrosit, genç", "Oligodendrosit, yaşlı", "OPC, genç", "OPC, yaşlı"],
    "Ortanca (log2)": [8.456, 8.717, 7.830, 7.808],
})
conv["CPM karşılığı"] = (2 ** conv["Ortanca (log2)"] - 1).round(0)
r = put(ws, r, "log2 değerlerinin CPM karşılığı (ortanca üzerinden)", conv,
        {"Ortanca (log2)": F3, "CPM karşılığı": INT})

ag = pd.read_csv(RESULTS / "Kif13a_AgingMouse_OPC_Oligo_age_summary.csv")
ag = ag[["donor_age_category", "subclass_name", "n_cells", "n_donors",
         "mean_Kif13a", "median_Kif13a", "fraction_expressing"]]
ag.columns = ["Yaş grubu", "Alt sınıf", "Hücre", "Fare", "Ortalama (log2)",
              "Ortanca (log2)", "İfade eden oran"]
ag["Yaş grubu"] = ag["Yaş grubu"].map({"adult": "genç", "aged": "yaşlı"})
put(ws, r, "Hücreler havuzlanmış özet", ag,
    {"Hücre": INT, "Ortalama (log2)": F3, "Ortanca (log2)": F3, "İfade eden oran": PCT})
autofit(ws)

# ------------------------------------------------------------- 3. donors
ws = wb.create_sheet("Yaşlanma donörler")
u = donors.drop_duplicates("donor_label").copy()


def to_days(a):
    a = str(a)
    if a.endswith("wks"):
        return int(a.split()[0]) * 7
    if a.endswith("M"):
        return int(a[:-1]) * 30
    return int(a[1:])


u["gün"] = u.donor_age.map(to_days)
adult, aged = u[u.donor_age_category == "adult"], u[u.donor_age_category == "aged"]
grp = pd.DataFrame({
    "Grup": ["Genç yetişkin", "Yaşlı"],
    "Fare sayısı": [len(adult), len(aged)],
    "En küçük yaş (gün)": [adult["gün"].min(), aged["gün"].min()],
    "En büyük yaş (gün)": [adult["gün"].max(), aged["gün"].max()],
    "Ortanca yaş (gün)": [adult["gün"].median(), aged["gün"].median()],
    "Dişi": [(adult.donor_sex == "F").sum(), (aged.donor_sex == "F").sum()],
    "Erkek": [(adult.donor_sex == "M").sum(), (aged.donor_sex == "M").sum()],
})
r = put(ws, 1, "Grupların özeti", grp)

for cat, label in [("adult", "Genç yetişkin grubu, yaşa göre fare sayısı"),
                   ("aged", "Yaşlı grup, yaşa göre fare sayısı")]:
    s = u[u.donor_age_category == cat]
    t = (s.groupby(["donor_age", "gün"])
           .agg(**{"Fare": ("donor_label", "nunique"),
                   "Dişi": ("donor_sex", lambda x: (x == "F").sum()),
                   "Erkek": ("donor_sex", lambda x: (x == "M").sum())})
           .reset_index().sort_values("gün")
           .rename(columns={"donor_age": "Yaş etiketi", "gün": "Gün"}))
    r = put(ws, r, label, t)

cells = (donors.groupby(["donor_age_category", "subclass_name"])
         .agg(**{"Fare": ("donor_label", "nunique"), "Toplam hücre": ("n_cells", "sum"),
                 "Fare başı ortanca": ("n_cells", "median"),
                 "En az": ("n_cells", "min"), "En çok": ("n_cells", "max")})
         .reset_index()
         .rename(columns={"donor_age_category": "Yaş grubu", "subclass_name": "Alt sınıf"}))
cells["Yaş grubu"] = cells["Yaş grubu"].map({"adult": "genç", "aged": "yaşlı"})
put(ws, r, "Fare başına hücre katkısı", cells,
    {"Toplam hücre": INT, "Fare başı ortanca": INT, "En az": INT, "En çok": INT})
autofit(ws)

# ------------------------------------------------------------- 4. regional
ws = wb.create_sheet("Bölgesel 13 parça")
w = pd.read_csv(RESULTS / "Kif13a_WMB-10Xv3_all_regions_OPC_Oligo_summary.csv")
opc = w[w.subclass == "326 OPC NN"].set_index("region")
oli = w[w.subclass == "327 Oligo NN"].set_index("region").reindex(opc.index)
reg = pd.DataFrame({
    "Parça": opc.index,
    "OPC ortalama": opc.mean_Kif13a.values,
    "Oligo ortalama": oli.mean_Kif13a.values,
    "Oligo - OPC": oli.mean_Kif13a.values - opc.mean_Kif13a.values,
    "OPC ifade oran": opc.fraction_expressing.values,
    "Oligo ifade oran": oli.fraction_expressing.values,
}).sort_values("OPC ortalama", ascending=False)
r = put(ws, 1, "13 parçada OPC ve olgun oligodendrosit", reg,
        {"OPC ortalama": F3, "Oligo ortalama": F3, "Oligo - OPC": "+0.000;-0.000",
         "OPC ifade oran": PCT, "Oligo ifade oran": PCT})

a_, b_ = reg["OPC ortalama"], reg["Oligo ortalama"]
fo, fl = reg["OPC ifade oran"], reg["Oligo ifade oran"]
spread = pd.DataFrame({
    "Ölçüt": ["En düşük", "En yüksek", "Aralık", "Standart sapma",
              "Değişim katsayısı (%)", "İfade eden oran aralığı"],
    "OPC": [a_.min(), a_.max(), a_.max() - a_.min(), a_.std(ddof=1),
            100 * a_.std(ddof=1) / a_.mean(), fo.max() - fo.min()],
    "Oligo": [b_.min(), b_.max(), b_.max() - b_.min(), b_.std(ddof=1),
              100 * b_.std(ddof=1) / b_.mean(), fl.max() - fl.min()],
})
spread["Oran (OPC/Oligo)"] = spread["OPC"] / spread["Oligo"]
r = put(ws, r, "Yayılım karşılaştırması", spread,
        {"OPC": F3, "Oligo": F3, "Oran (OPC/Oligo)": "0.00"})

tests = pd.DataFrame({
    "Test": ["Varyans oranı (F testi)", "Levene (Brown-Forsythe)",
             "Korteks+HPF vs arka beyin, OPC", "Korteks+HPF vs arka beyin, Oligo"],
    "İstatistik": [9.87, 24.95, None, None],
    "Fark (log2)": [None, None, 1.335, 0.107],
    "p değeri": [0.000372, 0.00003, 0.0179, 0.1964],
})
put(ws, r, "İstatistik testleri", tests,
    {"İstatistik": "0.00", "Fark (log2)": "+0.000;-0.000", "p değeri": "0.00000"})
autofit(ws)

# ------------------------------------------------------------- 5. aging by region
ws = wb.create_sheet("Yaşlanma bölge")
ar = pd.read_csv(RESULTS / "Kif13a_AgingMouse_OPC_Oligo_age_region_summary.csv")
ar = ar[["anatomical_division_label", "donor_age_category", "subclass_name",
         "n_cells", "n_donors", "mean_Kif13a", "median_Kif13a", "fraction_expressing"]]
ar.columns = ["Bölge", "Yaş grubu", "Alt sınıf", "Hücre", "Fare",
              "Ortalama (log2)", "Ortanca (log2)", "İfade eden oran"]
ar = ar.sort_values(["Alt sınıf", "Bölge", "Yaş grubu"])
ar["Yaş grubu"] = ar["Yaş grubu"].map({"adult": "genç", "aged": "yaşlı"})
put(ws, 1, "Yaşlanma kohortunun 7 bölgesi, yaş grubuna göre", ar,
    {"Hücre": INT, "Ortalama (log2)": F3, "Ortanca (log2)": F3, "İfade eden oran": PCT})
autofit(ws)

# ------------------------------------------------------------- 6. replication
ws = wb.create_sheet("Kohort tekrarı")
w2 = pd.read_csv(RESULTS / "Kif13a_WMB-10Xv3_all_regions_OPC_Oligo_summary.csv")
ag2 = pd.read_csv(RESULTS / "Kif13a_AgingMouse_OPC_Oligo_age_region_summary.csv")
w2["division"] = w2.region.replace({"Isocortex-1": "Isocortex", "Isocortex-2": "Isocortex"})
w2["prod"] = w2.mean_Kif13a * w2.n_cells
wm = w2.groupby(["division", "subclass"], as_index=False).agg(prod=("prod", "sum"), n=("n_cells", "sum"))
wm["ref"] = wm["prod"] / wm["n"]
adult_reg = (ag2[ag2.donor_age_category == "adult"]
             .rename(columns={"anatomical_division_label": "division",
                              "subclass_name": "subclass", "mean_Kif13a": "aging"})
             [["division", "subclass", "aging"]])
mg = wm.merge(adult_reg, on=["division", "subclass"])
r = 1
for sc, label, rho, pv in [("326 OPC NN", "OPC", 0.964, 0.0005),
                           ("327 Oligo NN", "Olgun oligodendrosit", 0.893, 0.0068)]:
    s = mg[mg.subclass == sc].sort_values("ref", ascending=False)
    t = pd.DataFrame({
        "Bölge": s.division.values,
        "Referans atlas (log2)": s.ref.values,
        "Yaşlanma kohortu (log2)": s.aging.values,
        "Fark": s.aging.values - s.ref.values,
    })
    r = put(ws, r, "{0} (Spearman rho = {1:+.3f}, p = {2})".format(label, rho, pv), t,
            {"Referans atlas (log2)": F3, "Yaşlanma kohortu (log2)": F3, "Fark": "+0.000;-0.000"})
autofit(ws)

# ------------------------------------------------------------- 7. whole brain
ws = wb.create_sheet("Tüm beyin sıralaması")
r = put(ws, 1, "En yüksek 10 hücre tipi", top10_df, {"Hücre": INT, "Ortalama (log2)": F3})
r = put(ws, r, "Glia tipleri karşılaştırması", glia_df,
        {"Ortalama (log2)": F3, "İfade eden oran": "0.00"})
r = put(ws, r, "En düşük hücreler", low_df,
        {"Ortalama (log2)": F3, "İfade eden oran": "0.00"})
full = rank[["Sıra", "subclass", "cell_class", "neurotransmitter", "n_cells",
             "mean_Kif13a", "fraction_expressing", "n_partitions"]].reset_index(drop=True)
full.columns = ["Sıra", "Alt sınıf", "Sınıf", "Nörotransmitter", "Hücre",
                "Ortalama (log2)", "İfade eden oran", "Parça sayısı"]
put(ws, r, "Tam sıralama, 265 alt sınıf", full,
    {"Hücre": INT, "Ortalama (log2)": F3, "İfade eden oran": PCT})
autofit(ws)

# ------------------------------------------------------------- 8. detection vs level
ws = wb.create_sheet("Tespit vs seviye")
dv = pd.read_csv(RESULTS / "derived_detection_vs_level.csv")
dv = dv[["subclass", "rank_overall", "n_cells", "mean_Kif13a", "fraction_expressing",
         "cpm_among_expressing", "naive_fold", "detection_ratio", "level_ratio_among_expressing"]]
dv.columns = ["Alt sınıf", "Genel sıra", "Hücre", "Ortalama (log2, sıfırlar dahil)",
              "Tespit oranı", "Tespit edilende seviye (CPM)", "Yanıltıcı kat",
              "Tespit oranı katı", "Seviye katı"]
put(ws, 1, "Tespit oranı ile ifade seviyesinin ayrılması", dv,
    {"Hücre": INT, "Ortalama (log2, sıfırlar dahil)": F3, "Tespit oranı": PCT,
     "Tespit edilende seviye (CPM)": INT, "Yanıltıcı kat": "0.0",
     "Tespit oranı katı": "0.00", "Seviye katı": "0.00"})
autofit(ws)

# ------------------------------------------------------------- 9. dataset
ws = wb.create_sheet("Veri seti")
ds = pd.DataFrame({
    "Özellik": ["Ad", "Hücre sayısı", "Parça / bölge", "Fare sayısı", "Gen sayısı",
                "İfade matrisi (disk)", "Yöntem", "Değer birimi", "Metadata sürümü"],
    "Referans atlas": ["WMB-10Xv3", "2.341.350", "13 parça", "çıkarılmadı", "32.285",
                       "yaklaşık 95 GB", "10x Genomics Chromium v3", "log2(CPM+1)", "20241115"],
    "Yaşlanma kohortu": ["Zeng Aging Mouse 10Xv3", "1.162.565", "7 anatomik bölge",
                         "108 (64 genç, 44 yaşlı)", "32.285", "yaklaşık 13 GB",
                         "10x Genomics Chromium v3", "log2(CPM+1)", "20250131"],
})
r = put(ws, 1, "Veri setleri", ds)

frames = []
for f in sorted(RESULTS.glob("Kif13a_WMB-10Xv3-*_subclass_summary.csv")):
    region = re.search(r"WMB-10Xv3-(.+)_subclass_summary", f.name).group(1)
    t = pd.read_csv(f, keep_default_na=False)
    t["region"] = region
    frames.append(t)
allp = pd.concat(frames, ignore_index=True)
part = (allp.groupby("region")
        .agg(**{"Hücre": ("n_cells", "sum"), "Alt sınıf": ("subclass", "nunique")})
        .reset_index().rename(columns={"region": "Parça"})
        .sort_values("Hücre", ascending=False))
part["Yüzde"] = part["Hücre"] / part["Hücre"].sum()
r = put(ws, r, "Parçalara göre hücre dağılımı", part, {"Hücre": INT, "Yüzde": PCT})

tax = pd.DataFrame({
    "Düzey": ["Sınıf (class)", "Alt sınıf (subclass)", "Süpertip (supertype)",
              "Küme (cluster)", "Nörotransmitter tipi"],
    "Sayı": [34, 338, 1201, 5322, 10],
})
put(ws, r, "Ortak taksonomi", tax, {"Sayı": INT})
autofit(ws)

wb.active = 0
wb.save(OUT)
print("Kaydedildi:", OUT)
print("Sayfalar:", wb.sheetnames)
